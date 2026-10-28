#!/usr/bin/env python3
"""Regenerates the synthetic part of the round-trip corpus.

Requires python-docx, openpyxl, python-pptx and Pillow (pip install python-docx openpyxl
python-pptx pillow) and, for the LibreOffice exports, `soffice` on the PATH.

    python3 test/corpus/tools/generate_corpus.py test/corpus

Files written by the libraries themselves are named `pylib-*`; files re-exported by LibreOffice
are named `libreoffice-*`. Files authored by Word, Excel and PowerPoint come from the Apache POI
test data and are not generated here (see test/corpus/README.md).
"""
import datetime
import io
import os
import shutil
import subprocess
import sys
import tempfile

from PIL import Image, ImageDraw

OUT = sys.argv[1] if len(sys.argv) > 1 else "test/corpus"


def png(w, h, color):
    img = Image.new("RGB", (w, h), color)
    d = ImageDraw.Draw(img)
    d.ellipse((w // 4, h // 4, 3 * w // 4, 3 * h // 4), fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def make_docx(path, pages=None):
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor, Inches

    d = docx.Document()
    d.core_properties.author = "Corpus generator"
    d.core_properties.created = datetime.datetime(2024, 1, 1)
    d.core_properties.modified = datetime.datetime(2024, 1, 1)
    d.core_properties.last_modified_by = "Corpus generator"
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Header text"
    sec.footer.paragraphs[0].text = "Footer text"
    if pages:
        for i in range(pages):
            d.add_heading(f"Chapter {i + 1}", level=1)
            for j in range(9):
                p = d.add_paragraph(
                    f"Page {i + 1} paragraph {j + 1}. " +
                    "The quick brown fox jumps over the lazy dog. " * 4)
                if j == 3:
                    p.runs[0].bold = True
            if i < pages - 1:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.save(path)
        return
    for level in range(1, 7):
        d.add_heading(f"Heading level {level}", level=level)
    p = d.add_paragraph("Plain, ")
    p.add_run("bold, ").bold = True
    p.add_run("italic, ").italic = True
    p.add_run("underlined, ").underline = True
    r = p.add_run("struck, ")
    r.font.strike = True
    r = p.add_run("red 16pt, ")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r = p.add_run("highlighted.")
    r.font.highlight_color = 7  # yellow
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.left_indent = Inches(0.5)
    for t in ["First bullet", "Second bullet"]:
        d.add_paragraph(t, style="List Bullet")
    d.add_paragraph("Nested bullet", style="List Bullet 2")
    for t in ["First number", "Second number"]:
        d.add_paragraph(t, style="List Number")
    table = d.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for i in range(3):
        for j in range(3):
            table.cell(i, j).text = f"R{i + 1}C{j + 1}"
    a = table.cell(2, 0).merge(table.cell(2, 1))
    a.text = "Merged"
    d.add_paragraph("Image below:")
    d.add_picture(png(200, 120, (30, 90, 200)), width=Inches(2))
    # Hyperlink built from raw XML (python-docx has no API for it).
    p = d.add_paragraph("Visit ")
    rel = d.part.relate_to("https://example.org/", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "example.org"
    run.append(t)
    h.append(run)
    p._p.append(h)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("After the page break.")
    new = d.add_section()
    new.orientation = WD_ORIENT.LANDSCAPE
    new.page_width, new.page_height = new.page_height, new.page_width
    d.add_paragraph("Landscape section.")
    d.save(path)


def make_xlsx(path, rows=None):
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    if rows:
        ws.append(["Id", "Name", "Amount", "Date", "Flag"])
        base = datetime.date(2024, 1, 1)
        for i in range(1, rows + 1):
            ws.append([i, f"Item {i}", round(i * 1.37, 2), base + datetime.timedelta(days=i % 365), i % 2 == 0])
        ws.cell(row=rows + 2, column=3, value=f"=SUM(C2:C{rows + 1})")
        wb.save(path)
        return
    ws.append(["Product", "Units", "Price", "Total", "Date", "In stock"])
    data = [("Apples", 10, 0.5), ("Pears", 4, 0.75), ("Plums", 25, 0.2), ("Kiwis", 7, 1.1)]
    for i, (name, units, price) in enumerate(data, start=2):
        ws.append([name, units, price, f"=B{i}*C{i}", datetime.date(2024, 3, i), units > 5])
        ws.cell(row=i, column=3).number_format = '"$"#,##0.00'
        ws.cell(row=i, column=4).number_format = '"$"#,##0.00'
        ws.cell(row=i, column=5).number_format = "yyyy-mm-dd"
    ws["A7"] = "Sum"
    ws["B7"] = "=SUM(B2:B5)"
    ws["B8"] = "=AVERAGE(B2:B5)"
    ws["B9"] = "=MIN(B2:B5)"
    ws["B10"] = "=MAX(B2:B5)"
    ws["B11"] = "=COUNT(B2:B5)"
    ws["B12"] = '=IF(B7>40,"many","few")'
    ws["B13"] = "=ROUND(C3*3.333,2)"
    ws["B14"] = '=VLOOKUP("Plums",A2:C5,3,FALSE)'
    ws["B15"] = '=CONCAT(A2," & ",A3)'
    ws["C16"] = 0.256
    ws["C16"].number_format = "0.0%"
    for c in ws[1]:
        c.font = Font(bold=True, italic=True)
        c.fill = PatternFill("solid", fgColor="DDEEFF")
        c.border = Border(bottom=Side(style="thin"))
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["A"].width = 18
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:F5"
    ws.conditional_formatting.add("B2:B5", CellIsRule(operator="greaterThan", formula=["8"], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.merge_cells("D7:E7")
    chart = BarChart()
    chart.title = "Units"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(chart, "H2")
    other = wb.create_sheet("Second sheet")
    other["A1"] = "Cross-sheet"
    other["A2"] = "=Data!B7*2"
    other["B1"] = True
    other["B2"] = "text with spaces  "
    wb.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Corpus deck"
    s.placeholders[1].text = "Generated for round-trip tests"
    s.notes_slide.notes_text_frame.text = "Speaker notes on the title slide"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Bullets"
    tf = s.placeholders[1].text_frame
    tf.text = "First point"
    for t, lvl in [("Second point", 0), ("Sub point", 1)]:
        p = tf.add_paragraph()
        p.text = t
        p.level = lvl
    run = tf.paragraphs[0].runs[0]
    run.font.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x20, 0x60, 0xA0)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Shapes and picture"
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2), Inches(2), Inches(1))
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(2), Inches(2), Inches(1))
    s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.5), Inches(2), Inches(2), Inches(1))
    s.shapes.add_picture(png(160, 120, (200, 60, 40)), Inches(1), Inches(4), Inches(2))
    tb = s.shapes.add_textbox(Inches(4), Inches(4), Inches(4), Inches(1))
    tb.text_frame.text = "A free text box"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Chart and table"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("Sales", (3, 5, 4))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(4), Inches(3), cd)
    t = s.shapes.add_table(2, 2, Inches(5), Inches(1.5), Inches(4), Inches(1)).table
    t.cell(0, 0).text = "A"
    t.cell(1, 1).text = "D"
    prs.save(path)


def libreoffice_export(src, dst):
    filters = {".docx": "docx:MS Word 2007 XML", ".xlsx": "xlsx:Calc MS Excel 2007 XML",
               ".pptx": "pptx:Impress MS PowerPoint 2007 XML"}
    ext = os.path.splitext(src)[1]
    with tempfile.TemporaryDirectory() as tmp:
        profile = "file://" + os.path.join(tmp, "profile")
        subprocess.run(["soffice", f"-env:UserInstallation={profile}", "--headless", "--convert-to",
                        filters[ext], "--outdir", tmp, src], check=True, capture_output=True, timeout=300)
        out = os.path.join(tmp, os.path.basename(src))
        shutil.move(out, dst)


def main():
    for sub in ("docx", "xlsx", "pptx", "large"):
        os.makedirs(os.path.join(OUT, sub), exist_ok=True)
    targets = [
        (make_docx, "docx/pylib-features.docx"),
        (make_xlsx, "xlsx/pylib-features.xlsx"),
        (make_pptx, "pptx/pylib-features.pptx"),
    ]
    for fn, rel in targets:
        path = os.path.join(OUT, rel)
        fn(path)
        lo = os.path.join(OUT, os.path.dirname(rel), "libreoffice-" + os.path.basename(rel).replace("pylib-", ""))
        if shutil.which("soffice"):
            libreoffice_export(path, lo)
    make_docx(os.path.join(OUT, "large/pylib-200-pages.docx"), pages=200)
    make_xlsx(os.path.join(OUT, "large/pylib-10000-rows.xlsx"), rows=10000)


if __name__ == "__main__":
    main()
